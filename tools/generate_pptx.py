#!/usr/bin/env python3
"""
Banki.ru Corporate PPTX Generator
Based on the official "Шаблон демо.pptx" template and "Цвета Банки.ру" theme.

Usage:
    python3 generate_pptx.py '<json_data>'
    python3 generate_pptx.py --file <json_file>
"""

import json
import sys
import os
from pptx import Presentation
from pptx.util import Cm, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# =============================================================================
# CORPORATE BRAND CONSTANTS — "Цвета Банки.ру" theme
# =============================================================================

# Theme colors (from a:clrScheme)
COLOR_ACCENT1 = RGBColor(0x0D, 0x8B, 0xFF)  # #0D8BFF - Primary Blue
COLOR_ACCENT2 = RGBColor(0x96, 0x41, 0xFF)  # #9641FF - Purple
COLOR_ACCENT3 = RGBColor(0x00, 0xD7, 0x3C)  # #00D73C - Green
COLOR_ACCENT4 = RGBColor(0xFF, 0x78, 0x28)  # #FF7828 - Orange
COLOR_ACCENT5 = RGBColor(0x96, 0xDC, 0xFF)  # #96DCFF - Light Blue (subtitles on dark bg)
COLOR_ACCENT6 = RGBColor(0xC8, 0xBE, 0xFF)  # #C8BEFF - Light Purple
COLOR_DK2 = RGBColor(0xEB, 0xF5, 0xFF)      # #EBF5FF - Ice Blue bg
COLOR_LT2 = RGBColor(0xF0, 0xF0, 0xFF)      # #F0F0FF - Lavender bg

# Extended palette (from template content)
COLOR_BLACK = RGBColor(0x00, 0x00, 0x00)         # #000000
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)          # #FFFFFF
COLOR_DARK_NAVY = RGBColor(0x0F, 0x1A, 0x24)     # #0F1A24
COLOR_DARK_SUBTITLE = RGBColor(0x16, 0x21, 0x36)  # #162136
COLOR_BODY_GRAY = RGBColor(0x56, 0x61, 0x73)     # #566173

# PRIORITY BACKGROUNDS (blue is default)
BG_BLUE = RGBColor(0x3E, 0x89, 0xF7)        # #3E89F7 - Blue section bg (DEFAULT)
BG_DARK = RGBColor(0x1E, 0x25, 0x37)        # #1E2537 - Dark navy section bg
BG_LIGHT = RGBColor(0xEB, 0xF5, 0xFF)       # #EBF5FF - Light ice blue bg
BG_PURPLE = RGBColor(0x96, 0x41, 0xFF)      # #9641FF - Purple section bg

# Font
FONT_NAME = "Coil Regular"
FONT_FALLBACK = "Arial"

# Slide dimensions (standard 16:9)
SLIDE_WIDTH = Cm(25.4)
SLIDE_HEIGHT = Cm(14.29)

# Layout constants
MARGIN_LEFT = Cm(1.06)
MARGIN_TOP = Cm(1.06)
TITLE_HEIGHT = Cm(1.91)
CONTENT_TOP = Cm(4.0)

# Logo paths
ASSETS_DIR = os.path.join(os.path.expanduser("~"), ".config", "opencode", "tools", "assets")
LOGO_WHITE = os.path.join(ASSETS_DIR, "bankiru_logo_white.png")  # For blue/dark backgrounds
LOGO_DARK = os.path.join(ASSETS_DIR, "bankiru_logo_dark.png")    # For light backgrounds
LOGO_SMALL = os.path.join(ASSETS_DIR, "bankiru_logo.png")        # Fallback

# Logo dimensions (standard format, proportionally scaled from demo)
LOGO_WIDTH_STD = Cm(4.8)
LOGO_HEIGHT_STD = Cm(1.46)
LOGO_LEFT = Cm(0.7)
LOGO_TOP = Cm(0.6)


def get_logo_path(on_dark_bg=True):
    """Return the appropriate logo file for the background."""
    if on_dark_bg:
        if os.path.exists(LOGO_WHITE):
            return LOGO_WHITE
    else:
        if os.path.exists(LOGO_DARK):
            return LOGO_DARK
    if os.path.exists(LOGO_SMALL):
        return LOGO_SMALL
    return None


def set_font(run, size_pt, color, bold=False):
    """Apply corporate font styling to a text run."""
    run.font.name = FONT_NAME
    run.font.size = Pt(size_pt)
    run.font.color.rgb = color
    run.font.bold = bold


def add_text_box(slide, left, top, width, height, text, size_pt, color,
                 alignment=PP_ALIGN.LEFT):
    """Add a styled text box to a slide."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    p = tf.paragraphs[0]
    p.alignment = alignment
    run = p.add_run()
    run.text = text
    set_font(run, size_pt, color)
    return txBox


def add_text_with_brand(slide, left, top, width, height, text, size_pt, color,
                        alignment=PP_ALIGN.LEFT):
    """Add text where 'Банки.ру' is highlighted in accent1 blue."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = alignment

    brand = "Банки.ру"
    parts = text.split(brand)

    for i, part in enumerate(parts):
        if part:
            run = p.add_run()
            run.text = part
            set_font(run, size_pt, color)
        if i < len(parts) - 1:
            run = p.add_run()
            run.text = brand
            set_font(run, size_pt, COLOR_ACCENT1)

    return txBox


def set_slide_bg(slide, color):
    """Set a solid background color for a slide."""
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_logo(slide, on_dark_bg=True):
    """Add the Banki.ru logo to a slide (top-left)."""
    logo_path = get_logo_path(on_dark_bg)
    if logo_path is None:
        return None
    return slide.shapes.add_picture(
        logo_path, LOGO_LEFT, LOGO_TOP, LOGO_WIDTH_STD, LOGO_HEIGHT_STD
    )


def add_rounded_rect(slide, left, top, width, height, fill_color=None):
    """Add a rounded rectangle shape."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    shape.line.fill.background()
    return shape


# =============================================================================
# SLIDE BUILDERS
# =============================================================================

def build_title_slide(prs, data):
    """Title/cover slide — blue background with white logo and text."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    set_slide_bg(slide, BG_BLUE)

    # Logo (white on blue bg)
    add_logo(slide, on_dark_bg=True)

    # Title
    title_text = data.get("title", "Презентация")
    add_text_box(
        slide, MARGIN_LEFT, Cm(4.5),
        Cm(17.0), Cm(5.0),
        title_text, 42.8, COLOR_WHITE
    )

    # Subtitle
    subtitle = data.get("subtitle", "")
    if subtitle:
        add_text_box(
            slide, MARGIN_LEFT, Cm(9.5),
            Cm(17.0), Cm(2.0),
            subtitle, 14.0, COLOR_ACCENT5  # #96DCFF light blue on colored bg
        )

    return slide


def build_section_slide(prs, data):
    """Section divider — blue background with white logo and text."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BG_BLUE)

    # Logo (white on blue bg)
    add_logo(slide, on_dark_bg=True)

    # Section title — lower portion
    title_text = data.get("title", "Раздел")
    add_text_box(
        slide, MARGIN_LEFT, Cm(7.0),
        Cm(20.0), Cm(3.5),
        title_text, 36.0, COLOR_WHITE
    )

    # Subtitle
    subtitle = data.get("subtitle", "")
    if subtitle:
        add_text_box(
            slide, MARGIN_LEFT, Cm(10.5),
            Cm(20.0), Cm(2.0),
            subtitle, 14.0, COLOR_ACCENT5  # #96DCFF
        )

    return slide


def build_content_slide(prs, data):
    """Standard content slide — white background."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Title
    title_text = data.get("title", "")
    if title_text:
        add_text_with_brand(
            slide, MARGIN_LEFT, MARGIN_TOP,
            Cm(22.0), TITLE_HEIGHT,
            title_text, 25.6, COLOR_BLACK
        )

    # Subtitle
    subtitle = data.get("subtitle", "")
    if subtitle:
        add_text_box(
            slide, MARGIN_LEFT, Cm(3.0),
            Cm(22.0), Cm(1.0),
            subtitle, 12.8, COLOR_DARK_SUBTITLE
        )

    # Bullets
    bullets = data.get("bullets", [])
    if bullets:
        start_top = Cm(4.5) if subtitle else CONTENT_TOP
        txBox = slide.shapes.add_textbox(
            Cm(1.5), start_top, Cm(22.0), Cm(9.0)
        )
        tf = txBox.text_frame
        tf.word_wrap = True

        for i, bullet in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_before = Pt(6)
            p.space_after = Pt(4)

            # Blue accent bar as bullet marker
            marker = p.add_run()
            marker.text = "\u2502  "  # Vertical bar character
            set_font(marker, 10.7, COLOR_ACCENT1)

            # Bullet text — handle brand name highlighting
            brand = "Банки.ру"
            if brand in bullet:
                parts = bullet.split(brand)
                for j, part in enumerate(parts):
                    if part:
                        run = p.add_run()
                        run.text = part
                        set_font(run, 10.7, COLOR_BODY_GRAY)
                    if j < len(parts) - 1:
                        run = p.add_run()
                        run.text = brand
                        set_font(run, 10.7, COLOR_ACCENT1)
            else:
                run = p.add_run()
                run.text = bullet
                set_font(run, 10.7, COLOR_BODY_GRAY)

    return slide


def build_two_column_slide(prs, data):
    """Two-column comparison — white background with blue card headers."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Title
    title_text = data.get("title", "")
    if title_text:
        add_text_with_brand(
            slide, MARGIN_LEFT, MARGIN_TOP,
            Cm(22.0), TITLE_HEIGHT,
            title_text, 25.6, COLOR_BLACK
        )

    col_width = Cm(11.0)
    col_gap = Cm(0.5)
    col_top = CONTENT_TOP
    card_height = Cm(9.5)

    # --- Left column ---
    left_title = data.get("left_title", "")
    left_bullets = data.get("left_bullets", [])

    # Card background
    add_rounded_rect(slide, MARGIN_LEFT, col_top, col_width, card_height)

    # Blue header strip
    add_rounded_rect(
        slide, MARGIN_LEFT, col_top,
        col_width, Cm(1.0),
        fill_color=COLOR_ACCENT1
    )

    if left_title:
        add_text_box(
            slide, Cm(1.46), Cm(4.15),
            Cm(10.0), Cm(0.8),
            left_title, 12.8, COLOR_WHITE
        )

    if left_bullets:
        txBox = slide.shapes.add_textbox(Cm(1.46), Cm(5.3), Cm(10.0), Cm(7.5))
        tf = txBox.text_frame
        tf.word_wrap = True
        for i, bullet in enumerate(left_bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_before = Pt(4)
            p.space_after = Pt(2)
            marker = p.add_run()
            marker.text = "\u2502  "
            set_font(marker, 9.0, COLOR_ACCENT1)
            run = p.add_run()
            run.text = bullet
            set_font(run, 9.0, COLOR_BODY_GRAY)

    # --- Right column ---
    right_left = Cm(1.06) + col_width + col_gap
    right_title = data.get("right_title", "")
    right_bullets = data.get("right_bullets", [])

    add_rounded_rect(slide, right_left, col_top, col_width, card_height)

    add_rounded_rect(
        slide, right_left, col_top,
        col_width, Cm(1.0),
        fill_color=COLOR_ACCENT1
    )

    if right_title:
        add_text_box(
            slide, right_left + Cm(0.4), Cm(4.15),
            Cm(10.0), Cm(0.8),
            right_title, 12.8, COLOR_WHITE
        )

    if right_bullets:
        txBox = slide.shapes.add_textbox(right_left + Cm(0.4), Cm(5.3), Cm(10.0), Cm(7.5))
        tf = txBox.text_frame
        tf.word_wrap = True
        for i, bullet in enumerate(right_bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_before = Pt(4)
            p.space_after = Pt(2)
            marker = p.add_run()
            marker.text = "\u2502  "
            set_font(marker, 9.0, COLOR_ACCENT1)
            run = p.add_run()
            run.text = bullet
            set_font(run, 9.0, COLOR_BODY_GRAY)

    return slide


def build_metrics_slide(prs, data):
    """Metrics/stats slide — white background with blue-accented metric cards."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Title
    title_text = data.get("title", "")
    if title_text:
        add_text_with_brand(
            slide, MARGIN_LEFT, MARGIN_TOP,
            Cm(22.0), TITLE_HEIGHT,
            title_text, 25.6, COLOR_BLACK
        )

    metrics = data.get("metrics", [])
    if not metrics:
        return slide

    count = len(metrics)
    cols = min(count, 4)
    card_width = Cm(5.5)
    card_height = Cm(4.2)
    gap_x = Cm(0.5)
    gap_y = Cm(0.5)
    start_left = MARGIN_LEFT
    start_top = CONTENT_TOP

    for idx, metric in enumerate(metrics):
        row = idx // cols
        col = idx % cols

        left = start_left + (card_width + gap_x) * col
        top = start_top + (card_height + gap_y) * row

        # Card body
        add_rounded_rect(slide, left, top, card_width, card_height)

        # Blue accent strip at top
        add_rounded_rect(
            slide, left, top,
            card_width, Cm(0.6),
            fill_color=COLOR_ACCENT1
        )

        # Value (large blue number)
        value = metric.get("value", "")
        add_text_box(
            slide, left + Cm(0.3), top + Cm(0.9),
            card_width - Cm(0.6), Cm(2.0),
            value, 33.9, COLOR_ACCENT1
        )

        # Label
        label = metric.get("label", "")
        add_text_box(
            slide, left + Cm(0.3), top + Cm(2.9),
            card_width - Cm(0.6), Cm(1.0),
            label, 10.2, COLOR_BODY_GRAY
        )

    return slide


def build_image_slide(prs, data):
    """Image slide — white background with image and caption."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Title
    title_text = data.get("title", "")
    if title_text:
        add_text_with_brand(
            slide, MARGIN_LEFT, MARGIN_TOP,
            Cm(22.0), TITLE_HEIGHT,
            title_text, 25.6, COLOR_BLACK
        )

    # Image or placeholder
    image_path = data.get("image_path", "")
    if image_path and os.path.exists(image_path):
        slide.shapes.add_picture(
            image_path, MARGIN_LEFT, CONTENT_TOP,
            Cm(22.0), Cm(8.5)
        )
    else:
        placeholder = add_rounded_rect(
            slide, MARGIN_LEFT, CONTENT_TOP,
            Cm(22.0), Cm(8.5),
            fill_color=COLOR_DK2  # #EBF5FF light blue tint
        )
        add_text_box(
            slide, Cm(8.0), Cm(7.5),
            Cm(10.0), Cm(2.0),
            "[Изображение]", 14.0, COLOR_BODY_GRAY,
            alignment=PP_ALIGN.CENTER
        )

    # Caption
    caption = data.get("caption", "")
    if caption:
        add_text_box(
            slide, MARGIN_LEFT, Cm(12.8),
            Cm(22.0), Cm(1.0),
            caption, 8.0, COLOR_DARK_SUBTITLE
        )

    return slide


def build_thank_you_slide(prs, data):
    """Final slide — blue background with white logo and thank-you text."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BG_BLUE)

    # Logo (white on blue bg)
    add_logo(slide, on_dark_bg=True)

    # Thank you text — large, lower-left
    title_text = data.get("title", "Спасибо за внимание!")
    add_text_box(
        slide, MARGIN_LEFT, Cm(7.5),
        Cm(15.0), Cm(4.5),
        title_text, 53.4, COLOR_WHITE
    )

    # Contact info
    contact = data.get("contact", "")
    if contact:
        add_text_box(
            slide, MARGIN_LEFT, Cm(12.0),
            Cm(15.0), Cm(1.5),
            contact, 12.0, COLOR_ACCENT5  # #96DCFF
        )

    return slide


# =============================================================================
# MAIN GENERATOR
# =============================================================================

SLIDE_BUILDERS = {
    "title": build_title_slide,
    "section": build_section_slide,
    "content": build_content_slide,
    "two_column": build_two_column_slide,
    "metrics": build_metrics_slide,
    "image": build_image_slide,
    "thank_you": build_thank_you_slide,
}


def generate_presentation(data):
    """Generate a complete PPTX presentation from JSON data."""
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    slides_data = data.get("slides", [])
    if not slides_data:
        raise ValueError("No slides defined in JSON data")

    for slide_data in slides_data:
        slide_type = slide_data.get("type", "content")
        builder = SLIDE_BUILDERS.get(slide_type, build_content_slide)
        builder(prs, slide_data)

    filename = data.get("filename", "presentation.pptx")
    out_dir = os.path.dirname(filename)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)

    prs.save(filename)
    return os.path.abspath(filename)


def main():
    if len(sys.argv) < 2:
        print("Usage: python3 generate_pptx.py '<json_data>'")
        print("       python3 generate_pptx.py --file <json_file>")
        sys.exit(1)

    if sys.argv[1] == "--file":
        if len(sys.argv) < 3:
            print("Error: --file requires a file path argument")
            sys.exit(1)
        with open(sys.argv[2], 'r', encoding='utf-8') as f:
            data = json.load(f)
    else:
        data = json.loads(sys.argv[1])

    output_path = generate_presentation(data)
    print(f"Presentation saved to: {output_path}")


if __name__ == "__main__":
    main()
